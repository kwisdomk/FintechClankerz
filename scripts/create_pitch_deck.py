from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = "FINTECH_AI_MVP_Pitch.pptx"
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY=RGBColor(10,20,37); PANEL=RGBColor(20,35,56); PANEL2=RGBColor(25,45,67); TEAL=RGBColor(58,220,190); WHITE=RGBColor(240,246,250); MUTED=RGBColor(164,183,201); GOLD=RGBColor(255,193,94)

def bg(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0, Inches(.12), Inches(7.5)); sh.fill.solid(); sh.fill.fore_color.rgb=TEAL; sh.line.fill.background()
def text(s,x,y,w,h,value,size=16,color=WHITE,bold=False,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.02);tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0];p.text=value;p.font.name="Aptos";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
    if align is not None:p.alignment=align
    return sh
def card(s,x,y,w,h,fill=PANEL):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=fill;sh.line.color.rgb=fill;return sh
def bullets(s,x,y,w,items,size=14,step=.48):
    for i,item in enumerate(items):
        text(s,x,y+i*step,.22,.3,"•",size+2,TEAL,True);text(s,x+.27,y+i*step,w-.27,step+.02,item,size,WHITE)
def slide(kicker,heading,sub=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);bg(s);text(s,.65,.32,12,.25,kicker.upper(),10,TEAL,True);text(s,.65,.68,12,.65,heading,28,WHITE,True)
    if sub:text(s,.67,1.35,12,.4,sub,13,MUTED)
    n=len(prs.slides);text(s,.68,7.15,10,.2,"FINTECH AI  /  DIGITAL SURFERS",9,MUTED);text(s,12.1,7.12,.5,.2,f"{n}/5",9,TEAL,True,align=PP_ALIGN.RIGHT)
    return s

# 1: problem + users
s=slide("FINANCIAL AUDIT MVP","Find the signal. Keep the auditor in control.","FINTECH AI — Digital Surfers | Evidence-led investigation of unusual financial transactions")
card(s,.75,2.05,5.75,4.45,PANEL);text(s,1.05,2.35,5,.3,"THE PROBLEM",13,GOLD,True)
bullets(s,1.05,2.95,5.05,["High-volume transaction spreadsheets make unusual items hard to spot.","A raw anomaly flag lacks context and follow-up evidence.","Review decisions and investigation history need a consistent record."],16,.88)
card(s,6.8,2.05,5.75,4.45,PANEL2);text(s,7.1,2.35,5,.3,"INTENDED USERS",13,TEAL,True)
bullets(s,7.1,2.95,5.05,["Internal / external auditors", "Finance control and compliance teams", "Audit leads who need a review trail"],16,.88)
text(s,1.05,5.95,10.9,.4,"AI assists investigation; it does not decide fraud or replace professional judgment.",14,TEAL,True)

# 2: proposed solution
s=slide("PROPOSED SOLUTION","Import → screen → explain → review","Deterministic data checks work without AI. Gemma is an optional evidence-grounded explanation layer.")
steps=[("01  IMPORT","Validate CSV rows and persist transactions."),("02  SCREEN","Flag repeated references and supplier + amount + date matches."),("03  EXPLAIN","Gemma interprets server-supplied evidence; no database/tool access."),("04  REVIEW","Auditor confirms or dismisses; decision and timestamp are stored.")]
for i,(h,d) in enumerate(steps):
    x=.78+i*3.15;card(s,x,2.25,2.85,2.05,PANEL if i%2==0 else PANEL2);text(s,x+.2,2.52,2.5,.35,h,14,TEAL,True);text(s,x+.2,3.08,2.45,.95,d,14,WHITE)
    if i<3:text(s,x+2.86,3.05,.28,.35,"›",20,TEAL,True,align=PP_ALIGN.CENTER)
card(s,.78,4.75,11.8,1.55,RGBColor(23,53,62));text(s,1.05,5.0,11,.3,"EVIDENCE BOUNDARY",12,TEAL,True)
text(s,1.05,5.38,11,.65,"Observed evidence is persisted transaction / rule output. Gemma contributes interpretation and recommended checks—not new transaction facts or financial actions.",15,WHITE)

# 3: implementation + architecture
s=slide("IMPLEMENTATION","The repo contains the complete MVP stack.","Connected now: frontend → backend → SQLite → local Ollama/Gemma. The broader analysis pipeline is included; API hookup is still partial.")
cols=[(.75,"FRONTEND",["Next.js 16 / React","Dashboard, upload, registry, findings"]),(3.9,"API + DB",["Python / FastAPI / Pydantic","CSV validation + REST","SQLite audit records"]),(7.05,"DATA + AI",["Pandas / NumPy / Isolation Forest pipeline","Deterministic duplicate rules in API","Ollama • Gemma 4 local"]),(10.2,"DEMO ACCESS",["Cloudflare Quick Tunnel","Backend + Ollama proxy paths","Synthetic data only"])]
for x,h,lines in cols:
    card(s,x,2.15,2.35,2.6,PANEL);text(s,x+.17,2.42,2,.3,h,12,TEAL,True);bullets(s,x+.17,2.92,2.02,lines,12,.53)
for x in [3.25,6.4,9.55]:text(s,x,3.15,.4,.35,"→",19,TEAL,True,align=PP_ALIGN.CENTER)
card(s,.75,5.08,11.8,1.3,PANEL2);text(s,1.02,5.28,11,.3,"REPRODUCIBLE EVIDENCE",12,GOLD,True)
text(s,1.02,5.68,11.2,.5,"100 synthetic rows accepted • pipeline reports 6 duplicate rows + 34 anomaly rows • Gemma 4 7.5B Q4_K_M • local short-prompt benchmark ≈7.1 generated tokens/sec",13,WHITE)

# 4: live demonstration + impact
s=slide("LIVE DEMO & IMPACT","Follow one finding from data to human decision.","Temporary demo URL — the link can change when the Quick Tunnel restarts.")
card(s,.75,2.05,6.1,4.4,PANEL);text(s,1.05,2.35,5.4,.3,"LIVE ENDPOINTS",12,TEAL,True)
bullets(s,1.05,2.9,5.4,["App: https://skating-peak-academics-simulation.trycloudflare.com", "Backend health: /backend/api/health", "Swagger: /backend/docs", "Ollama: POST /ollama/api/generate", "Findings: /api/findings → /explain → /review"],13,.63)
card(s,7.15,2.05,5.4,4.4,PANEL2);text(s,7.48,2.35,4.8,.3,"PROJECT IMPACT",12,GOLD,True)
bullets(s,7.48,2.9,4.75,["Reduce time spent searching spreadsheets.","Keep evidence, interpretation, and auditor decisions distinct.","Local inference keeps demo records on-device.","Impact target: measure triage time and review completeness with auditors."],14,.77)
text(s,1.05,6.02,11,.25,"The public demo uses synthetic records; the tunnel and model run on the developer machine.",11,MUTED)

# 5: open-source + DigitalOcean deployment research
s=slide("OPEN SOURCE + CLOUD ROADMAP","Public code is the OSS deliverable; cloud inference is the next deployment step.","Current deployment is local Ollama. DigitalOcean hosting is researched/planned, not provisioned or integrated yet.")
card(s,.75,2.0,5.75,4.6,PANEL);text(s,1.05,2.3,5,.3,"OPEN-SOURCE EVIDENCE",12,TEAL,True)
bullets(s,1.05,2.82,5.05,["Public repo: github.com/kwisdomk/FintechClankerz", "Main + mbr-3 include frontend, API, data pipeline, sample CSV and reports.", "README documents local run and API.", "Open-source license file is still needed; agree license with contributors."],14,.78)
card(s,6.8,2.0,5.75,4.6,PANEL2);text(s,7.1,2.3,5,.3,"DIGITALOCEAN INFERENCE PLAN",12,GOLD,True)
bullets(s,7.1,2.82,5.05,["GPU Droplet runs Ollama + Gemma 4; keep port 11434 private.","API points OLLAMA_BASE_URL to private GPU IP; expose only HTTPS app/API.","RTX 4000 Ada (20GB) is a plausible first test for 6.6GB Q4 model weights; benchmark context/memory.","For occasional use, start GPU for a scheduled window; persist model volume and verify compute/storage billing before stopping."],12,.7)
text(s,.82,6.72,11.8,.23,"DO GPU Droplets: digitalocean.com/products/gpu-droplets (listed options include RTX 4000 Ada, RTX 6000 Ada, H100/H200; pricing/region availability changes). Inference engine: Ollama.",9,MUTED)

prs.core_properties.title="FINTECH AI — Digital Surfers: Financial Audit Investigation MVP"
prs.core_properties.subject="Five-slide MVP: problem, solution, implementation, demo, impact, open-source evidence, DigitalOcean inference roadmap"
prs.core_properties.author="FINTECH AI / Digital Surfers"
prs.save(OUT)
print(OUT)
